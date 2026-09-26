"""Le canevas du personnel à remplir, et le fichier déposé qu'on retélécharge : relus tels quels par l'import."""
import io
from datetime import date

import openpyxl

from courtage.fichier import lire_fichier
from courtage.fichier.canevas import canevas
from tests.outils import V1, deposer, en_tant_que, fichier_azito


def test_le_canevas_rempli_se_relit_sans_rien_retoucher():
    wb = openpyxl.load_workbook(io.BytesIO(canevas()))
    assert wb.sheetnames[0] == "Personnel"                        # l'import lit la première feuille
    assert {"Mode d'emploi", "Exemple"} <= set(wb.sheetnames)
    ws = wb["Personnel"]
    entete = next(r for r in ws.iter_rows(values_only=True) if r and r[0] == "Matricule")
    assert not any("nom" in str(c).lower() for c in entete if c)  # pas de colonne de nom à remplir
    ligne = [r.row for r in ws["A"] if r.value == "Matricule"][0] + 1
    for i, (m, n, e, s, c, x) in enumerate([("A001", date(1970, 5, 2), date(1998, 3, 1), 450_000, "Cadre", "F"),
                                            ("A002", date(1985, 1, 20), date(2012, 9, 15), 210_000, "Employé", "M")]):
        for j, v in enumerate((m, n, e, s, c, x)):
            ws.cell(row=ligne + i, column=j + 1, value=v)
    tampon = io.BytesIO()
    wb.save(tampon)
    lecture = lire_fichier(tampon.getvalue(), "personnel.xlsx")
    assert [a for a in lecture.anomalies if a.niveau == "bloquant"] == []
    assert lecture.periodicite == "mensuel"
    assert [(l.matricule, l.salaire_annuel, l.categorie, l.sexe) for l in lecture.lignes] == [
        ("A001", 5_400_000, "Cadre", "F"), ("A002", 2_520_000, "Employé", "M")]


def test_un_canevas_vide_ne_passe_pas_pour_un_personnel():
    lecture = lire_fichier(canevas(), "canevas.xlsx")
    assert any(a.code == "fichier_vide" for a in lecture.anomalies)


def test_le_canevas_se_telecharge_sans_compte(client):
    r = client.get(f"{V1}/referentiel/canevas-personnel")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "canevas-personnel.xlsx" in r.headers["content-disposition"]


def test_un_fichier_depose_se_retelecharge_et_se_recharge_a_l_identique(client, azito):
    r = client.get(f"{V1}/organisations/{azito['org']}/fichiers/{azito['fichier']}/telechargement",
                   headers=en_tant_que(azito["drh"]))
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[:2] == ["Personnel", "Anomalies"]
    brut = " ".join(str(c) for row in wb["Personnel"].iter_rows(values_only=True) for c in row if c)
    assert "Nom" not in brut                                     # un nom n'a jamais été gardé
    relu = lire_fichier(r.content, "retour.xlsx")
    original = lire_fichier(fichier_azito(), "azito.xlsx")
    cle = lambda l: (l.matricule, l.naissance, l.embauche, l.salaire_annuel, l.categorie)  # noqa: E731
    assert [cle(l) for l in relu.lignes] == [cle(l) for l in original.lignes]
    # Et il se redépose : le même personnel.
    assert deposer(client, azito["org"], azito["drh"], r.content)["effectif"] == 23
    # Un étranger au dossier ne le télécharge pas.
    etranger = client.get(f"{V1}/organisations/{azito['org']}/fichiers/{azito['fichier']}/telechargement",
                          headers=en_tant_que(azito["etranger"]))
    assert etranger.status_code == 403
