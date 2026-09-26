"""Les exports Excel : l'étude et les réponses des assureurs, lisibles et recalculables."""
import io

import openpyxl
import pytest

from tests.outils import V1, en_tant_que, etude
from tests.test_rapport import emettre
from tests.test_reponses import cahier, repondre  # noqa: F401  (la fixture « cahier »)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def classeur(r):
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(XLSX) and "attachment" in r.headers["content-disposition"]
    return openpyxl.load_workbook(io.BytesIO(r.content))


def valeurs(ws):
    return {row[0]: row[1] for row in ws.iter_rows(values_only=True) if row and row[0]}


def test_l_etude_en_excel(client, azito):
    e = emettre(client, azito)
    lue = client.get(f"{V1}/organisations/{azito['org']}/etudes/{e['id']}", headers=en_tant_que(azito["drh"])).json()
    wb = classeur(client.get(f"{V1}/organisations/{azito['org']}/etudes/{e['id']}/export", headers=en_tant_que(azito["drh"])))
    assert wb.sheetnames == ["Synthèse", "Échéancier", "Par catégorie", "Sensibilités", "Salariés"]
    synthese = valeurs(wb["Synthèse"])
    assert synthese["Dette actuarielle (F CFA)"] == lue["totaux"]["dette"]
    assert synthese["Cotisation totale (F CFA)"] == lue["totaux"]["cotisation_totale"]
    assert synthese["Rapport scellé"] == lue["rapport"]["numero"]
    assert "Taux d'actualisation" in synthese
    # Le contrôle se recalcule : une formule, pas un chiffre recopié.
    assert str(synthese["Contrôle : dette + charge − fonds (F CFA)"]).startswith("=MAX(")

    ech = wb["Échéancier"]
    entete = next(r for r in ech.iter_rows(values_only=True) if r and r[0] == "Année")
    assert entete[:7] == ("Année", "Départs", "Si tous partent (F CFA)", "Prestations probables (F CFA)",
                          "Valeur actuelle (F CFA)", "Cumul des probables (F CFA)", "Part du total")
    lignes = [r for r in ech.iter_rows() if isinstance(r[0].value, int)]
    assert [r[0].value for r in lignes] == [a["annee"] for a in lue["echeancier"]]
    assert str(lignes[1][5].value).startswith("=") and str(lignes[0][6].value).startswith("=")
    total = next(r for r in ech.iter_rows() if r[0].value == "Total")
    assert str(total[3].value).startswith("=SUM(")

    salaries = [r for r in wb["Salariés"].iter_rows(values_only=True) if r and r[0] and r[0] not in ("Matricule", "Total")]
    assert len([r for r in salaries if isinstance(r[1], (str, type(None)))]) >= 23
    brut = " ".join(str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c)
    assert "Nom" not in brut.replace("Nombre", "")                      # aucun nom de salarié : il n'y en a pas
    assert all(c.font.name == "Arial" for row in wb["Synthèse"].iter_rows() for c in row if c.value is not None)


def test_un_brouillon_s_exporte_et_le_dit(client, azito):
    e = etude(client, azito).json()
    wb = classeur(client.get(f"{V1}/organisations/{azito['org']}/etudes/{e['id']}/export", headers=en_tant_que(azito["drh"])))
    assert valeurs(wb["Synthèse"])["Statut"] == "Brouillon : pas encore émise, ces chiffres peuvent changer"


def test_les_reponses_en_excel(client, cahier):  # noqa: F811
    repondre(client, cahier)
    repondre(client, cahier, assureur="Assureur B", transfert_penalite=0.05, taux_garanti=0.035)
    url = f"{V1}/organisations/{cahier['org']}/fiches/{cahier['fiche']['id']}/reponses"
    lu = client.get(url, headers=en_tant_que(cahier["drh"])).json()
    wb = classeur(client.get(f"{url}/export", headers=en_tant_que(cahier["drh"])))
    assert wb.sheetnames == ["Comparaison", "Conformité", "Scénarios", "Cahier des charges"]
    comp = wb["Comparaison"]
    entete = next(r for r in comp.iter_rows(values_only=True) if r and r[0] == "Rang")
    rangs = [r for r in comp.iter_rows(values_only=True) if r and isinstance(r[0], int)]
    assert [r[1] for r in rangs] == [x["assureur"] for x in lu["reponses"]]
    recommandee = next(x["assureur"] for x in lu["reponses"] if x["id"] == lu["recommandee"])
    assert [r[entete.index("Recommandée")] for r in rangs if r[1] == recommandee] == ["oui"]
    conf = [r for r in wb["Conformité"].iter_rows(values_only=True) if r and r[0] == "Assureur B"]
    assert any(r[1] == "Pénalité de transfert" and r[4] == "non" for r in conf)
    scen = [r for r in wb["Scénarios"].iter_rows(values_only=True) if r and r[0] in ("Assureur A", "Assureur B")]
    assert {r[1] for r in scen} == {"prudent", "central", "favorable"}


def test_un_etranger_n_exporte_rien(client, azito):
    e = emettre(client, azito)
    r = client.get(f"{V1}/organisations/{azito['org']}/etudes/{e['id']}/export", headers=en_tant_que(azito["etranger"]))
    assert r.status_code == 403
