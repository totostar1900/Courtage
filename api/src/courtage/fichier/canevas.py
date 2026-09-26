"""Le canevas du personnel à remplir, et le fichier déposé qu'on retélécharge.

Les deux ont la même feuille « Personnel », en PREMIÈRE position parce que l'import lit la première feuille :
un fichier téléchargé se corrige et se redépose tel quel. Aucune colonne de nom : la plateforme n'en lit
jamais (spec §2.8), et n'en a donc aucun à rendre. Module pur : des octets xlsx, sans base.
"""
import io
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ENCRE, DISCRET, FOND = "141A33", "6A7192", "EEF0F7"
_TITRE = Font(bold=True, size=13, color=ENCRE)
_ENTETE = Font(bold=True, color="FFFFFF")
_FOND_ENTETE = PatternFill("solid", fgColor="2F419A")
_NOTE = Font(italic=True, color=DISCRET)

COLONNES = [
    # intitulé (reconnu par l'import), largeur, format, obligatoire, ce qu'il faut y mettre, exemple
    ("Matricule", 14, "@", True, "L'identifiant du salarié dans votre paie. Jamais son nom.", "D014"),
    ("Date de naissance", 18, "DD/MM/YYYY", True, "Au format jj/mm/aaaa.", "02/05/1970"),
    ("Date d'embauche", 18, "DD/MM/YYYY", True, "L'entrée dans l'entreprise (ancienneté reprise comprise), jj/mm/aaaa.",
     "01/03/1998"),
    (None, 22, "#,##0", True, "Le salaire brut de base, en francs CFA, sans les primes exceptionnelles.", "450 000"),
    ("Catégorie", 16, "@", False, "Le nom de la catégorie telle que votre régime la nomme (Cadre, Employé…). "
     "Laisser vide si le régime est le même pour tous.", "Cadre"),
    ("Sexe", 8, "@", False, "F ou M. Facultatif.", "F"),
]
REGLES = [
    "Une ligne par salarié présent à la date d'arrêté, sans ligne vide entre deux salariés.",
    "Aucun nom, aucun prénom : le matricule suffit. Une colonne de nom serait ignorée sans être lue.",
    "Les dates au format jj/mm/aaaa ; les montants en chiffres, espaces permis (450 000).",
    "Ne pas renommer les intitulés de la ligne d'en-tête : c'est par eux que la plateforme reconnaît les colonnes.",
    "La feuille « Personnel » doit rester la première du classeur.",
    "À la dépose, vous indiquerez la date d'arrêté des données (en général une clôture : 31/12).",
]
EXEMPLE = [
    ("D001", date(1968, 4, 12), date(1994, 2, 1), 1_350_000, "Cadre", "M"),
    ("D002", date(1979, 11, 3), date(2006, 9, 18), 420_000, "Employé", "F"),
    ("D003", date(1990, 7, 25), date(2017, 1, 9), 285_000, "Employé", "M"),
]


def _intitule_salaire(periodicite: str) -> str:
    return "Salaire brut mensuel" if periodicite == "mensuel" else "Salaire brut annuel"


def _feuille_personnel(ws, titre: str, periodicite: str, lignes: list[tuple]) -> None:
    ws.title = "Personnel"
    ws.append([titre])
    ws["A1"].font = _TITRE
    ws.append(["Aucun nom : le matricule suffit. Une ligne par salarié."])
    ws["A2"].font = _NOTE
    entete = [c[0] or _intitule_salaire(periodicite) for c in COLONNES]
    ws.append(entete)
    for j, (_, largeur, _, obligatoire, _, _) in enumerate(COLONNES, start=1):
        cellule = ws.cell(row=3, column=j)
        cellule.font, cellule.fill = _ENTETE, _FOND_ENTETE
        cellule.alignment = Alignment(vertical="center")
        ws.column_dimensions[cellule.column_letter].width = largeur
    for ligne in lignes:
        ws.append(list(ligne))
    derniere = max(ws.max_row, 3 + 500)
    for j, (_, _, format_, _, _, _) in enumerate(COLONNES, start=1):
        for i in range(4, derniere + 1):
            ws.cell(row=i, column=j).number_format = format_
    ws.freeze_panes = "A4"
    zone = lambda col: f"{col}4:{col}{derniere}"  # noqa: E731
    dates = DataValidation(type="date", operator="between", formula1="DATE(1930,1,1)", formula2="DATE(2100,12,31)",
                           allow_blank=True, error="Une date au format jj/mm/aaaa.", showErrorMessage=True)
    dates.add(zone("B")); dates.add(zone("C"))
    salaire = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                             error="Un montant en francs CFA, positif.", showErrorMessage=True)
    salaire.add(zone("D"))
    sexe = DataValidation(type="list", formula1='"F,M"', allow_blank=True, error="F ou M.", showErrorMessage=True)
    sexe.add(zone("F"))
    for v in (dates, salaire, sexe):
        ws.add_data_validation(v)


def canevas() -> bytes:
    """Le classeur à remplir : « Personnel » (vide), « Mode d'emploi », « Exemple »."""
    wb = openpyxl.Workbook()
    _feuille_personnel(wb.active, "État du personnel au __/__/____", "mensuel", [])

    aide = wb.create_sheet("Mode d'emploi")
    aide.append(["Remplir l'état du personnel"])
    aide["A1"].font = _TITRE
    aide.append([])
    for regle in REGLES:
        aide.append([f"• {regle}"])
    aide.append([])
    aide.append(["Colonne", "Obligatoire", "Ce qu'il faut y mettre", "Exemple"])
    for cellule in aide[aide.max_row]:
        cellule.font, cellule.fill = _ENTETE, _FOND_ENTETE
    for intitule, _, _, obligatoire, explication, exemple in COLONNES:
        aide.append([intitule or _intitule_salaire("mensuel"), "oui" if obligatoire else "non", explication, exemple])
    aide.append([])
    aide.append(["Un salaire annuel ? Renommer la colonne « Salaire brut annuel » : la plateforme le reconnaît."])
    aide[f"A{aide.max_row}"].font = _NOTE
    for col, largeur in zip("ABCD", (24, 12, 80, 14)):
        aide.column_dimensions[col].width = largeur
    for ligne in aide.iter_rows(min_row=3):
        for cellule in ligne:
            cellule.alignment = Alignment(wrap_text=True, vertical="top")

    exemple = wb.create_sheet("Exemple")
    _feuille_personnel(exemple, "Exemple (fictif) : ne pas déposer cette feuille", "mensuel", EXEMPLE)
    exemple.title = "Exemple"
    return _octets(wb)


def exporter(lignes: list[dict], anomalies: list[dict], *, date_donnees: date, nom_fichier: str) -> bytes:
    """Le fichier déposé, tel que la plateforme le garde : les salaires annuels, sans nom, et ses anomalies."""
    wb = openpyxl.Workbook()
    valeurs = [(l["matricule"],
                date.fromisoformat(l["naissance"]) if l.get("naissance") else None,
                date.fromisoformat(l["embauche"]) if l.get("embauche") else None,
                l.get("salaire_annuel"), l.get("categorie"), l.get("sexe")) for l in lignes]
    _feuille_personnel(wb.active, f"État du personnel au {date_donnees.strftime('%d/%m/%Y')} — déposé sous "
                                  f"« {nom_fichier} »", "annuel", valeurs)
    ws = wb.create_sheet("Anomalies")
    ws.append(["Anomalies relevées à la dépose"])
    ws["A1"].font = _TITRE
    ws.append(["Niveau", "Ligne", "Colonne", "Message"])
    for cellule in ws[2]:
        cellule.font, cellule.fill = _ENTETE, _FOND_ENTETE
    for a in anomalies:
        ws.append([a.get("niveau"), a.get("ligne"), a.get("colonne"), a.get("message")])
    if not anomalies:
        ws.append(["—", None, None, "Aucune anomalie."])
    for col, largeur in zip("ABCD", (16, 8, 18, 90)):
        ws.column_dimensions[col].width = largeur
    return _octets(wb)


def _octets(wb) -> bytes:
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()
