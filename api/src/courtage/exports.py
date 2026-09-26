"""Les exports Excel : l'étude, et les réponses des assureurs. Module pur : un dict en clair, des octets xlsx.

Les résultats de l'étude sont des VALEURS : ceux du moteur, les mêmes que le rapport scellé, et la feuille le
dit. Ce qui s'en déduit est une FORMULE (cumuls, parts, totaux, contrôle du rapprochement, écarts des
sensibilités) : le classeur se recalcule si l'on y touche. Aucun nom de salarié : il n'y en a jamais eu.
"""
import io
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from courtage.services import hypotheses
from courtage.services.hypotheses import SENSIBILITES

POLICE = "Arial"
F = '#,##0;-#,##0;"-"'                 # francs CFA, entiers
PCT = '0.0%;-0.0%;"-"'
PCT2 = '0.00%;-0.00%;"-"'
_TITRE = Font(name=POLICE, bold=True, size=13, color="141A33")
_ENTETE = Font(name=POLICE, bold=True, color="FFFFFF")
_FOND = PatternFill("solid", fgColor="2F419A")
_GRAS = Font(name=POLICE, bold=True)
_NOTE = Font(name=POLICE, italic=True, color="6A7192")
_TEXTE = Font(name=POLICE)

HYPOTHESES = {
    "taux_actualisation": ("Taux d'actualisation", PCT2), "croissance_salaires": ("Croissance des salaires", PCT2),
    "inflation": ("Inflation", PCT2), "age_retraite": ("Âge de départ à la retraite", "0"),
    "taux_turnover": ("Taux de rotation du personnel", PCT2), "frais_sur_cotisation": ("Frais sur cotisation", PCT2),
    "table": ("Table de mortalité", "@"),
}
SCENARIOS = {"prudent": "prudent", "central": "central", "favorable": "favorable"}


def _jj(iso: str | None) -> str:
    return date.fromisoformat(iso[:10]).strftime("%d/%m/%Y") if iso else "—"


def _titre(ws, texte: str, note: str | None = None) -> None:
    ws.append([texte])
    ws.cell(row=ws.max_row, column=1).font = _TITRE
    if note:
        ws.append([note])
        ws.cell(row=ws.max_row, column=1).font = _NOTE
    ws.append([])


def _entete(ws, colonnes: list[str]) -> int:
    ws.append(colonnes)
    for c in ws[ws.max_row]:
        c.font, c.fill = _ENTETE, _FOND
        c.alignment = Alignment(wrap_text=True, vertical="center")
    return ws.max_row


def _largeurs(ws, largeurs: list[int]) -> None:
    for i, l in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = l


def _formats(ws, ligne_debut: int, ligne_fin: int, formats: dict[int, str]) -> None:
    for col, fmt in formats.items():
        for i in range(ligne_debut, ligne_fin + 1):
            ws.cell(row=i, column=col).number_format = fmt


def _police(wb) -> None:
    """Arial partout, en gardant le gras, la couleur et la taille déjà posés."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != POLICE:
                    c.font = Font(name=POLICE, bold=c.font.bold, italic=c.font.italic, color=c.font.color,
                                  size=c.font.size)


def _octets(wb) -> bytes:
    _police(wb)
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


# --- L'étude ------------------------------------------------------------------------------------

def etude(e: dict, organisation: str) -> bytes:
    """`e` : l'étude telle que `services.etudes.en_clair` la rend."""
    wb = openpyxl.Workbook()
    ref_dette = _synthese(wb.active, e, organisation)
    _echeancier(wb.create_sheet("Échéancier"), e)
    _par_categorie(wb.create_sheet("Par catégorie"), e)
    _sensibilites(wb.create_sheet("Sensibilités"), e, ref_dette)
    _salaries(wb.create_sheet("Salariés"), e)
    return _octets(wb)


def _synthese(ws, e: dict, organisation: str) -> str:
    ws.title = "Synthèse"
    _titre(ws, f"Étude IFC — {organisation} — au {_jj(e['date_evaluation'])}",
           "Résultats du moteur, les mêmes que le rapport scellé ; les lignes de contrôle sont des formules.")
    t, h = e["totaux"], e.get("hypotheses") or {}
    statut = ("Émise le " + _jj(e.get("emise_le"))) if e["statut"] == "emise" \
        else "Brouillon : pas encore émise, ces chiffres peuvent changer"
    regime = e.get("regime")
    lignes: list[tuple] = [
        ("Statut", statut, None),
        ("Rapport scellé", (e.get("rapport") or {}).get("numero") or "—", None),
        ("Convention collective", f"{e['convention']['libelle']} ({e['convention']['code']})", None),
        ("Régime", f"{regime['nom']}, version {regime['numero']}" if regime else "La convention seule", None),
        ("Moteur de calcul", e.get("version_moteur") or "—", None),
        ("Référentiel", e.get("referentiel_version") or "—", None),
    ]
    for libelle, valeur, fmt in lignes:
        ws.append([libelle, valeur])
    ws.append([])
    debut = _entete(ws, ["Totaux", "Valeur"]) + 1
    totaux = [("Effectif", t["effectif"], "0"), ("VAPF (F CFA)", t["vapf"], F),
              ("Dette actuarielle (F CFA)", t["dette"], F), ("Charge annuelle (F CFA)", t["charge"], F),
              ("Fonds constitué (F CFA)", e.get("fonds_disponible") or 0, F),
              ("Cotisation nette (F CFA)", t["cotisation_nette"], F),
              ("Cotisation totale (F CFA)", t["cotisation_totale"], F)]
    for libelle, valeur, fmt in totaux:
        ws.append([libelle, valeur])
        ws.cell(row=ws.max_row, column=2).number_format = fmt
    ligne = {libelle: debut + i for i, (libelle, _, _) in enumerate(totaux)}
    d, c, f = (f"B{ligne['Dette actuarielle (F CFA)']}", f"B{ligne['Charge annuelle (F CFA)']}",
               f"B{ligne['Fonds constitué (F CFA)']}")
    ws.append(["Contrôle : dette + charge − fonds (F CFA)", f"=MAX({d}+{c}-{f},0)"])
    ws.cell(row=ws.max_row, column=2).number_format = F
    ws.append(["Écart d'arrondi avec la cotisation nette (F CFA)",
               f"=B{ws.max_row}-B{ligne['Cotisation nette (F CFA)']}"])
    ws.cell(row=ws.max_row, column=2).number_format = F
    ws.append(["Le moteur calcule sans arrondi intermédiaire : l'écart d'arrondi vaut au plus quelques francs."])
    ws.cell(row=ws.max_row, column=1).font = _NOTE
    ws.append([])
    _entete(ws, ["Hypothèses", "Valeur"])
    for cle, (libelle, fmt) in HYPOTHESES.items():
        if cle in (h.get("valeurs") or {}):
            ws.append([libelle, h["valeurs"][cle]])
            ws.cell(row=ws.max_row, column=2).number_format = fmt
    if (h.get("valeurs") or {}).get("rotation_par_age"):
        ws.append(["Rotation par tranche d'âge (remplace le taux unique)",
                   hypotheses.en_texte("rotation_par_age", h["valeurs"]["rotation_par_age"])])
    for ecart in hypotheses.ecarts_en_texte(h.get("ecarts") or []):
        ws.append([f"Écart au référentiel : {ecart['libelle']}", f"{ecart['referentiel']} → {ecart['retenu']}"])
    if h.get("justification"):
        ws.append(["Justification", h["justification"]])
    _largeurs(ws, [48, 58])
    return f"'Synthèse'!{d}"


def _echeancier(ws, e: dict) -> None:
    _titre(ws, "Départs prévus, année par année",
           "Prestations probables : l'indemnité pondérée par la survie et la présence. Cumul et part : des formules.")
    annees = e["echeancier"]
    categories = sorted({c for a in annees for c in (a.get("par_categorie") or {})})
    libelle = lambda c: "Tout le personnel" if c == "*" and len(categories) == 1 else ("Autres salariés" if c == "*" else c)  # noqa: E731
    colonnes = ["Année", "Départs", "Si tous partent (F CFA)", "Prestations probables (F CFA)", "Valeur actuelle (F CFA)",
                "Cumul des probables (F CFA)", "Part du total"]
    colonnes += [f"Probables : {libelle(c)} (F CFA)" for c in categories]
    entete = _entete(ws, colonnes)
    premiere, derniere = entete + 1, entete + len(annees)
    total = derniere + 1
    for i, a in enumerate(annees):
        r = premiere + i
        cumul = f"=D{r}" if i == 0 else f"=F{r - 1}+D{r}"
        part = f"=IF($D${total}=0,0,D{r}/$D${total})"
        ws.append([a["annee"], a["effectif"], a["ifc"], a.get("prestations_probables", a["ifc"]), a["vapf"], cumul, part]
                  + [(a.get("par_categorie") or {}).get(c, {}).get("prestations_probables", 0) for c in categories])
    somme = lambda col: f"=SUM({col}{premiere}:{col}{derniere})"  # noqa: E731
    ws.append(["Total", somme("B"), somme("C"), somme("D"), somme("E"), f"=F{derniere}" if annees else 0,
               somme("G")] + [somme(get_column_letter(8 + j)) for j in range(len(categories))])
    for c in ws[total]:
        c.font = _GRAS
    formats = {2: "0", 3: F, 4: F, 5: F, 6: F, 7: PCT} | {8 + j: F for j in range(len(categories))}
    _formats(ws, premiere, total, formats)
    ws.freeze_panes = ws.cell(row=premiere, column=2)
    _largeurs(ws, [10, 10, 20, 22, 20, 22, 12] + [24] * len(categories))


def _par_categorie(ws, e: dict) -> None:
    _titre(ws, "Résultats par catégorie de personnel")
    entete = _entete(ws, ["Catégorie", "Effectif", "VAPF (F CFA)", "Dette (F CFA)", "Charge (F CFA)"])
    categories = e.get("par_categorie") or {}
    for c, v in sorted(categories.items()):
        ws.append(["Tout le personnel" if c == "*" and len(categories) == 1 else ("Autres salariés" if c == "*" else c),
                   v["effectif"], v["vapf"], v["dette"], v["charge"]])
    premiere, derniere = entete + 1, entete + max(len(categories), 1)
    ws.append(["Total"] + [f"=SUM({col}{premiere}:{col}{derniere})" for col in "BCDE"])
    for c in ws[ws.max_row]:
        c.font = _GRAS
    _formats(ws, premiere, ws.max_row, {2: "0", 3: F, 4: F, 5: F})
    _largeurs(ws, [24, 10, 20, 20, 20])


def _sensibilites(ws, e: dict, ref_dette: str) -> None:
    _titre(ws, "Sensibilités : la dette si une hypothèse bouge", "L'écart se calcule contre la dette de la synthèse.")
    entete = _entete(ws, ["Hypothèse modifiée", "Dette (F CFA)", "Écart à la dette de l'étude", "Charge (F CFA)"])
    ordre = list(SENSIBILITES)
    rangees = sorted((e.get("sensibilites") or {}).items(), key=lambda kv: ordre.index(kv[0]) if kv[0] in ordre else 99)
    for i, (cle, v) in enumerate(rangees):
        r = entete + 1 + i
        ws.append([SENSIBILITES.get(cle, cle), v["dette"], f"=IF({ref_dette}=0,0,B{r}/{ref_dette}-1)", v.get("charge")])
    _formats(ws, entete + 1, ws.max_row, {2: F, 3: PCT, 4: F})
    _largeurs(ws, [36, 20, 26, 20])


def _salaries(ws, e: dict) -> None:
    _titre(ws, "Le calcul, salarié par salarié", "Par matricule : aucun nom n'est lu ni gardé par la plateforme.")
    entete = _entete(ws, ["Matricule", "Catégorie", "Âge", "Ancienneté (ans)", "Date de retraite", "Mois dus",
                          "IFC à la retraite (F CFA)", "VAPF (F CFA)", "Dette (F CFA)", "Charge (F CFA)",
                          "Plancher de la convention appliqué", "Au-delà de l'âge de retraite"])
    lignes = e.get("lignes") or []
    for l in lignes:
        ws.append([l["matricule"], l.get("categorie"), l["age"], l["anciennete"], date.fromisoformat(l["date_retraite"]),
                   l.get("mois"), l["ifc"], l["vapf"], l["dette"], l["charge"],
                   "oui" if l.get("plancher_applique") else "non", "oui" if l.get("au_dela_de_la_retraite") else "non"])
    premiere, derniere = entete + 1, entete + max(len(lignes), 1)
    ws.append(["Total", None, None, None, None, None] + [f"=SUM({c}{premiere}:{c}{derniere})" for c in "GHIJ"])
    for c in ws[ws.max_row]:
        c.font = _GRAS
    _formats(ws, premiere, ws.max_row, {3: "0.00", 4: "0.00", 5: "DD/MM/YYYY", 6: "0.00", 7: F, 8: F, 9: F, 10: F})
    ws.freeze_panes = ws.cell(row=premiere, column=2)
    ws.auto_filter.ref = f"A{entete}:L{derniere}"
    _largeurs(ws, [12, 16, 8, 14, 14, 10, 20, 16, 16, 16, 16, 16])


# --- Les réponses des assureurs ---------------------------------------------------------------

def _oui(v) -> str:
    return "oui" if v is True else "non" if v is False else "—"


def reponses(t: dict, organisation: str) -> bytes:
    """`t` : les réponses telles que `services.reponses.tout` les rend."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Comparaison"
    _titre(ws, f"Réponses des assureurs — {organisation}",
           "Classées par coût net actualisé dans le scénario central ; la recommandée est la moins chère des conformes.")
    choix = t.get("choix")
    colonnes = ["Rang", "Assureur", "Reçue le", "Conforme au cahier", "Coût net actualisé (F CFA)", "Taux garanti",
                "Participation aux bénéfices", "Frais sur cotisations", "Frais sur encours", "Délai de paiement (jours)",
                "Préavis de transfert (mois)", "Pénalité de transfert", "Recommandée", "Choisie"]
    entete = _entete(ws, colonnes)
    for r in t["reponses"]:
        ws.append([r.get("rang"), r["assureur"], _jj(r.get("recue_le")), _oui(r.get("conforme")), r.get("cout_net_actualise"),
                   r.get("taux_garanti"), r.get("participation_benefices"), r.get("frais_sur_cotisations"),
                   r.get("frais_sur_encours"), r.get("delai_paiement_jours"), r.get("transfert_preavis_mois"),
                   r.get("transfert_penalite"), _oui(r["id"] == t.get("recommandee")),
                   _oui(bool(choix) and r["id"] == choix["reponse_id"])])
    _formats(ws, entete + 1, ws.max_row, {5: F, 6: PCT2, 7: PCT, 8: PCT2, 9: PCT2, 12: PCT2})
    if choix:
        ws.append([])
        ws.append(["Choix de l'entreprise", f"{choix['assureur']}, le {_jj(choix['choisi_le'])}"])
        if choix.get("motif"):
            ws.append(["Motif", choix["motif"]])
    _largeurs(ws, [8, 22, 12, 12, 22, 12, 14, 14, 14, 14, 14, 14, 13, 10])

    conf = wb.create_sheet("Conformité")
    _titre(conf, "Chaque réponse, critère par critère")
    entete = _entete(conf, ["Assureur", "Critère", "Demandé", "Offert", "Conforme"])
    for r in t["reponses"]:
        for c in r.get("conformite") or []:
            demande = _oui(c["demande"]) if isinstance(c["demande"], bool) else c["demande"]
            offert = _oui(c["offert"]) if isinstance(c["offert"], bool) or c["offert"] is None else c["offert"]
            conf.append([r["assureur"], c["libelle"], demande, offert, _oui(c["conforme"])])
    _largeurs(conf, [22, 38, 12, 12, 11])

    scen = wb.create_sheet("Scénarios")
    _titre(scen, "La projection du fonds, offre par offre et scénario par scénario",
           "Coût net actualisé : cotisations et découverts actualisés, moins le fonds restant à l'horizon.")
    entete = _entete(scen, ["Assureur", "Scénario", "Rendement", "Coût net actualisé (F CFA)", "Coût total (F CFA)",
                            "Frais totaux (F CFA)", "Fonds à l'horizon (F CFA)", "Couverture des départs restants",
                            "Années de découvert"])
    for o in ((t.get("comparaison") or {}).get("offres") or []):
        for s in o["scenarios"]:
            scen.append([o["nom"], SCENARIOS.get(s["scenario"], s["scenario"]), s["rendement"], round(s["cout_net_actualise"]),
                         round(s["cout_total"]), round(s["frais_totaux"]), round(s["fonds_final"]),
                         s.get("couverture_des_departs_restants"),
                         ", ".join(str(a) for a in s.get("annees_decouvert") or []) or "aucune"])
    _formats(scen, entete + 1, scen.max_row, {3: PCT, 4: F, 5: F, 6: F, 7: F, 8: PCT})
    _largeurs(scen, [22, 11, 11, 22, 18, 18, 20, 18, 22])

    cahier = wb.create_sheet("Cahier des charges")
    _titre(cahier, "Les conditions demandées", f"Réponses attendues avant le {_jj(t.get('date_limite_reponse'))}.")
    _entete(cahier, ["Condition", "Valeur"])
    for cle, valeur in (t.get("conditions") or {}).items():
        cahier.append([cle.replace("_", " "), _oui(valeur) if isinstance(valeur, bool) else valeur])
    _largeurs(cahier, [40, 16])
    return _octets(wb)
